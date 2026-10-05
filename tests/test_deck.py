"""The outline stays buildable: every slide parses, references existing figures and keeps the deck rules."""

import re

import pytest
from pptx import Presentation

from intent.data import load_config, resolve
from intent.deck import DIAGRAMS, SlideSpec, _md_tables, build_slide, parse_outline

CFG = load_config()
SPECS = parse_outline(resolve(CFG["deck"]["outline"]))
LAYOUTS = {4, 6, 8, 10}  # the layouts build_slide knows how to fill


def test_eleven_main_and_three_appendix_slides() -> None:
    keys = [s.key for s in SPECS]
    assert keys == [str(i) for i in range(1, 12)] + ["A1", "A2", "A3"]


@pytest.mark.parametrize("spec", SPECS, ids=lambda s: s.key)
def test_slide_is_complete_and_follows_the_rules(spec: SlideSpec) -> None:
    assert spec.title and spec.layout in LAYOUTS
    assert len(spec.lines) <= 3
    assert 1 <= len(re.findall(r"[.!?](?:\s|$)", spec.notes)) <= 3  # speaker notes: at most 3 sentences
    assert spec.figures or spec.tables or spec.diagram or spec.layout == 4
    assert spec.diagram is None or spec.diagram in DIAGRAMS
    for name in spec.figures:
        assert (resolve(CFG["deck"]["figures_dir"]) / name).exists(), name
    text = " ".join([spec.title, spec.notes, *spec.lines, *(c for t in spec.tables for r in t for c in r)])
    assert "—" not in text and "–" not in text  # no em or en dashes


def test_markdown_tables_drop_the_separator_row() -> None:
    block = "| a | b |\n|---|---|\n| 1 | 2 |\n\ntext\n| c |\n|:--:|\n| 3 |"
    assert _md_tables(block) == [[["a", "b"], ["1", "2"]], [["c"], ["3"]]]


def test_two_column_captions_go_under_figures_with_room() -> None:
    """Slide 8: the left figure fills its column, so both text lines go under the right one."""
    prs = Presentation(resolve(CFG["deck"]["template"]))
    spec = next(s for s in SPECS if s.key == "8")
    build_slide(prs, spec, CFG, "{72833802-FEF1-4C79-8D5D-14CF1EAF98D9}")
    slide = prs.slides[-1]
    captions = [sh for sh in slide.placeholders if sh.has_text_frame and sh.placeholder_format.idx in (1, 13)]
    assert [len(c.text_frame.paragraphs) for c in captions] == [2]
    assert captions[0].left > prs.slide_width // 2


def test_konfio_palette_is_the_one_requested() -> None:
    from intent.plot_style import KONFIO

    assert (KONFIO.surface, KONFIO.primary, KONFIO.secondary, KONFIO.neutral, KONFIO.warm, KONFIO.ink,
            KONFIO.grid) == ("#140527", "#8644F9", "#AD82FB", "#8A8499", "#D9637B", "#F2EEF8", "#2E2245")
