"""Spanish interview deck: template inspection, deck build, PDF export and PNG render.

Why inspect in code: the deck reuses only the template's layouts and style, never its
sample content, so the facts that matter (layout indices, placeholder geometry, theme
colors, fonts, the master's background) are read from the .pptx itself and written to
reports/template_inspection.md. plot_style.py sizes figures from these placeholder sizes.

Why build from the outline: reports/slides_outline_es.md is the single source of the deck's
words. The build parses it (title, layout, figures, markdown tables, native diagram, text
lines, speaker notes) instead of repeating the text here, so editing the outline and
rebuilding is the only way the slides change.

    python -m intent.deck inspect | build | pdf | png | check | all
"""

from __future__ import annotations

import copy
import re
import shutil
import subprocess
import sys
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from lxml import etree
from matplotlib import font_manager
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.shapes.placeholder import PicturePlaceholder, TablePlaceholder
from pptx.util import Emu, Inches, Pt

from intent import plot_style as ps
from intent.data import load_config, resolve

EMU_PER_INCH = 914400
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
THEME_SLOTS = ["dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5",
               "accent6", "hlink", "folHlink"]


def _inches(emu: int | None) -> str:
    return "-" if emu is None else f"{emu / EMU_PER_INCH:.2f}"


def _theme(prs) -> etree._Element:
    master = prs.slide_masters[0]
    return etree.fromstring(master.part.part_related_by(RT.THEME).blob)


def theme_colors(prs) -> list[tuple[str, str]]:
    """(slot, hex) for the 12 theme color slots; system colors resolve to their last value."""
    scheme = _theme(prs).find(".//a:clrScheme", NS)
    out = []
    for slot in THEME_SLOTS:
        node = scheme.find(f"a:{slot}", NS)[0]
        out.append((slot, "#" + (node.get("val") if node.tag.endswith("srgbClr") else node.get("lastClr"))))
    return out


def theme_fonts(prs) -> dict[str, str]:
    scheme = _theme(prs).find(".//a:fontScheme", NS)
    return {"name": scheme.get("name"),
            "major": scheme.find("a:majorFont/a:latin", NS).get("typeface"),
            "minor": scheme.find("a:minorFont/a:latin", NS).get("typeface")}


def master_facts(prs) -> dict[str, Any]:
    """Background reference, color map and default text sizes of the slide master."""
    root = prs.slide_masters[0].element
    bg = root.find(".//p:bg//a:schemeClr", NS)
    clr_map = root.find("p:clrMap", NS).attrib
    sizes = {}
    for style, path in [("title", "p:txStyles/p:titleStyle/a:lvl1pPr"),
                        ("body level 1", "p:txStyles/p:bodyStyle/a:lvl1pPr"),
                        ("body level 2", "p:txStyles/p:bodyStyle/a:lvl2pPr")]:
        rpr = root.find(f"{path}/a:defRPr", NS)
        sizes[style] = int(rpr.get("sz")) / 100
    return {"background": bg.get("val") if bg is not None else None, "clr_map": dict(clr_map), "sizes": sizes}


def _installed(font: str) -> bool:
    return any(f.name == font for f in font_manager.fontManager.ttflist)


def layouts(prs) -> list[dict[str, Any]]:
    out = []
    for i, layout in enumerate(prs.slide_layouts):
        phs = [{"idx": ph.placeholder_format.idx, "type": ph.placeholder_format.type.name,
                "left": _inches(ph.left), "top": _inches(ph.top),
                "width": _inches(ph.width), "height": _inches(ph.height)} for ph in layout.placeholders]
        deco = [f"{s.name} at ({_inches(s.left)}, {_inches(s.top)}), {_inches(s.width)} x {_inches(s.height)}"
                for s in layout.shapes if not s.is_placeholder]
        out.append({"index": i, "name": layout.name.strip(), "placeholders": phs, "decorations": deco})
    return out


def _content_placeholders(layout: dict[str, Any]) -> str:
    skip = {"DATE", "FOOTER", "SLIDE_NUMBER"}
    return "; ".join(f"{p['type'].lower()} idx {p['idx']} ({p['width']} x {p['height']})"
                     for p in layout["placeholders"] if p["type"] not in skip)


def inspection_markdown(cfg: dict[str, Any] | None = None) -> str:
    cfg = cfg or load_config()
    path = cfg["deck"]["template"]
    prs = Presentation(resolve(path))
    colors, fonts, master = theme_colors(prs), theme_fonts(prs), master_facts(prs)
    lays = layouts(prs)
    bg_slot = master["clr_map"].get(master["background"], master["background"])
    text_slot = master["clr_map"]["tx1"]
    hexes = dict(colors)

    md = [
        "# Template inspection",
        "",
        f"Source: `{path}`, read with python-pptx by `python -m intent.deck inspect`. The template's own "
        f"{len(prs.slides)} sample slides (text and photos) are unrelated to this project: the deck uses only "
        "its layouts, theme colors and fonts.",
        "",
        "## Slide and master",
        "",
        f"- Slide size: {_inches(prs.slide_width)} x {_inches(prs.slide_height)} in (16:9).",
        f"- Background: the master fills every layout with `{master['background']}`, which the color map sends "
        f"to `{bg_slot}` = **{hexes[bg_slot]}** (dark navy). No layout overrides it.",
        f"- Default text color: `tx1` maps to `{text_slot}` = **{hexes[text_slot]}**. The theme is dark: "
        "figures must use light ink on the navy, not dark ink on white.",
        "- Default sizes: " + ", ".join(f"{k} {v:g} pt" for k, v in master["sizes"].items()) + ".",
        "",
        "## Theme colors",
        "",
        f"Scheme name: {fonts['name']}.",
        "",
        "| Slot | Hex | Role in the template |",
        "|---|---|---|",
    ]
    roles = {"dk1": "unused by the master (bg1 maps here)", "lt1": "all text (tx1)",
             "dk2": "slide background (bg2)", "lt2": "secondary text (tx2)",
             "accent1": "green", "accent2": "teal; table header fill in the sample slides",
             "accent3": "magenta", "accent4": "rose", "accent5": "lavender gray", "accent6": "deep plum",
             "hlink": "hyperlinks", "folHlink": "visited hyperlinks"}
    md += [f"| {slot} | `{hx}` | {roles[slot]} |" for slot, hx in colors]
    md += [
        "",
        "## Fonts",
        "",
        "| Role | Theme font | Installed here | Used in figures |",
        "|---|---|---|---|",
        f"| Headings (major) | {fonts['major']} | {'yes' if _installed(fonts['major']) else 'no (Office cloud font)'} "
        f"| none: figures carry no titles; closest installed Didone is {ps.FONT_HEADING[0]} |",
        f"| Body (minor) | {fonts['minor']} | {'yes' if _installed(fonts['minor']) else 'no'} | every label |",
        "",
        "Gill Sans MT draws the digit 1 as a bare stroke, so `11.5%` can read as `II.5%`. Figures keep it for "
        "consistency with the slide text; the config hash on the traceability figure is set in Consolas.",
        "",
        "## Layouts",
        "",
        "Positions and sizes in inches (left, top; width x height). Every layout except 0 and 7 also has date "
        "(idx 10), footer (idx 11) and slide number (idx 12) placeholders at the bottom edge.",
        "",
        "| Index | Name | Content placeholders | Decorations (keep clear) |",
        "|---|---|---|---|",
    ]
    for lay in lays:
        deco = "<br>".join(lay["decorations"]) or "none"
        md.append(f"| {lay['index']} | {lay['name']} | {_content_placeholders(lay)} | {deco} |")
    md += ["", "### Placeholder geometry", ""]
    for lay in lays:
        md.append(f"**{lay['index']}. {lay['name']}**")
        md.append("")
        md.append("| idx | Type | Left | Top | Width | Height |")
        md.append("|---|---|---|---|---|---|")
        md += [f"| {p['idx']} | {p['type']} | {p['left']} | {p['top']} | {p['width']} | {p['height']} |"
               for p in lay["placeholders"]]
        md.append("")
    md += [
        "## Sample slides (to be removed in the deck)",
        "",
        "| # | Layout |",
        "|---|---|",
    ]
    md += [f"| {i + 1} | {s.slide_layout.name.strip()} |" for i, s in enumerate(prs.slides)]
    md += [
        "",
        "## How the deck uses this",
        "",
        "- `src/intent/plot_style.py` draws figures on the master background "
        f"(`{ps.SURFACE}`) with `{ps.INK}` and `{ps.INK_2}` ink and sizes them to the placeholders above: "
        f"full width {ps.FULL[0]} x {ps.FULL[1]} (layout 10), wide {ps.WIDE[0]} x {ps.WIDE[1]} (layout 8, "
        f"right area), half {ps.HALF[0]} x {ps.HALF[1]} (layout 6, each column).",
        "- Chart colors, in fixed order: "
        + ", ".join(f"`{c}`" for c in ps.CATEGORICAL)
        + f" (teal, rose, magenta, green), plus `{ps.NEUTRAL}` for de-emphasis. They are the theme accents, "
        "snapped to the dark-surface colorblind checks; the reasoning and the validator results are in the "
        "module docstring.",
        "- Layouts chosen per slide: `reports/slides_outline_es.md`.",
        "",
    ]
    return "\n".join(md)


def write_inspection(cfg: dict[str, Any] | None = None) -> str:
    cfg = cfg or load_config()
    out = resolve(cfg["deck"]["template_inspection"])
    out.write_text(inspection_markdown(cfg), encoding="utf-8")
    return str(out)


# ----------------------------------------------------------------------------- outline
@dataclass
class SlideSpec:
    """One section of the outline, already split into the parts the build places."""

    key: str                      # "1".."11", "A1".."A3"
    title: str
    layout: int
    figures: list[str] = field(default_factory=list)
    tables: list[list[list[str]]] = field(default_factory=list)
    diagram: str | None = None    # name of a native diagram ("pipeline", "despliegue")
    tag: str | None = None        # short label shown under the screenshot
    lines: list[str] = field(default_factory=list)
    notes: str = ""

    @property
    def appendix(self) -> bool:
        return self.key.startswith("A")


def _field(section: str, name: str) -> str:
    """Body of one top-level '- **Name:**' bullet, up to the next top-level bullet."""
    m = re.search(rf"^- \*\*{name}:\*\*(.*?)(?=^- \*\*|\Z)", section, flags=re.S | re.M)
    return m.group(1) if m else ""


def _md_tables(block: str) -> list[list[list[str]]]:
    """Markdown tables in a block, as lists of rows of cell strings (separator rows dropped)."""
    tables, current = [], []
    for line in block.splitlines() + [""]:
        s = line.strip()
        if s.startswith("|"):
            cells = [c.strip() for c in s.strip("|").split("|")]
            if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                current.append(cells)
        elif current:
            tables.append(current)
            current = []
    return tables


def parse_outline(path: Path) -> list[SlideSpec]:
    specs = []
    for section in re.split(r"\n## ", path.read_text(encoding="utf-8"))[1:]:
        head = re.match(r"(A?\d+)\. ", section)
        if head is None:  # "Figuras del deck" and other non-slide sections
            continue
        visual = _field(section, "Visual")
        tag = re.search(r'marca "([^"]+)"', visual)
        diagram = re.search(r"diagrama nativo `(\w+)`", visual)
        specs.append(SlideSpec(
            key=head.group(1),
            title=_field(section, "Título").strip(),
            layout=int(re.match(r"\s*(\d+)", _field(section, "Layout")).group(1)),
            figures=re.findall(r"`(\w+\.png)`", visual),
            tables=_md_tables(visual),
            diagram=diagram.group(1) if diagram else None,
            tag=tag.group(1) if tag else None,
            lines=[ln.strip()[2:] for ln in _field(section, "Texto").splitlines() if ln.strip().startswith("- ")],
            notes=_field(section, "Notas del orador").strip(),
        ))
    return specs


# ----------------------------------------------------------------------------- text helpers
TITLE_PT, TITLE_APPENDIX_PT = 28, 24  # the master's 48 pt fits a word, not a full-sentence title
BODY_PT, CAPTION_PT = 18, 16
LANG = "es-MX"                        # proofing language on every run


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#"))


def _plain(text: str) -> str:
    return text.replace("`", "")  # code spans in the outline are file or command names


def _runs(paragraph, text: str, size: float, color: str | None = None, bold: bool = False,
          font: str | None = None) -> None:
    """Add runs to a paragraph; **double asterisks** mark bold spans, as in the outline."""
    for part in re.split(r"(\*\*.+?\*\*)", _plain(text)):
        if not part:
            continue
        strong = part.startswith("**") and part.endswith("**")
        run = paragraph.add_run()
        run.text = part[2:-2] if strong else part
        run.font.size = Pt(size)
        if bold or strong:
            run.font.bold = True
        if color:
            run.font.color.rgb = _rgb(color)
        if font:
            run.font.name = font
        run._r.get_or_add_rPr().set("lang", LANG)


def _no_bullet(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.set("marL", "0")
    ppr.set("indent", "0")
    for old in ppr.findall(qn("a:buChar")) + ppr.findall(qn("a:buNone")):
        ppr.remove(old)
    etree.SubElement(ppr, qn("a:buNone"))


def _paragraphs(tf, lines: list[str], size: float, bullets: bool = True, color: str | None = None,
                align=None, font: str | None = None) -> None:
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        numbered = re.match(r"\d+\. ", line) is not None  # "1. ..." carries its own marker
        if not bullets or numbered:
            _no_bullet(p)
        if align is not None:
            p.alignment = align
        _runs(p, line, size, color=color, font=font)


def _remove(shape) -> None:
    shape._element.getparent().remove(shape._element)


def _box(ph) -> tuple[int, int, int, int]:
    """A placeholder's geometry in EMU (inherited from the layout when the slide sets none)."""
    return ph.left, ph.top, ph.width, ph.height


# ----------------------------------------------------------------------------- pictures and tables
def _fit(path: Path, box: tuple[int, int, int, int], native_dpi: int | None) -> tuple[int, int, int, int]:
    """Largest size that keeps the aspect ratio inside box, top aligned and centered horizontally.

    Figures drawn for the deck are never enlarged past their native size (native_dpi), so a
    point in the figure stays a point on the slide; screenshots (native_dpi None) just fit.
    """
    px_w, px_h = Image.open(path).size
    left, top, bw, bh = box
    w, h = bw, int(bw * px_h / px_w)
    if h > bh:
        w, h = int(bh * px_w / px_h), bh
    if native_dpi:
        nw = Inches(px_w / native_dpi)
        if w > nw:
            w, h = nw, int(nw * px_h / px_w)
    return left + (bw - w) // 2, top, w, h


def _picture(slide, path: Path, box, ph=None, native_dpi: int | None = None):
    """Picture inside box; into the placeholder ph when given, so it keeps the layout link."""
    left, top, w, h = _fit(path, box, native_dpi)
    if ph is None:
        return slide.shapes.add_picture(str(path), left, top, w, h)
    pic = PicturePlaceholder(ph._element, ph._parent).insert_picture(str(path))
    pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0  # insert_picture crops to fill
    pic.left, pic.top, pic.width, pic.height = left, top, w, h
    return pic


NUMERIC = re.compile(r"[\d\s.,\[\]+%-]+")


def _table(slide, rows: list[list[str]], box, ph, style_id: str) -> None:
    """Native table in the template's own table style, sized by its content.

    Font size follows the row count so the 24-row ablation still fits; column widths
    follow the longest text in each column. Rows are minimum heights: PowerPoint and
    LibreOffice grow a row when its text wraps.
    """
    n_rows, n_cols = len(rows), len(rows[0])
    size = 16 if n_rows <= 6 else 14 if n_rows <= 12 else 9.5
    frame = TablePlaceholder(ph._element, ph._parent).insert_table(n_rows, n_cols)
    left, top, width, _ = box
    tbl = frame.table
    tbl._tbl.tblPr.find(qn("a:tableStyleId")).text = style_id
    tbl.first_row, tbl.horz_banding = True, True
    plain = [[re.sub(r"\*\*", "", c) for c in row] for row in rows]
    # bold headers run ~25% wider than body text; a floor of 7 keeps "Hamming" or "PR-AUC" on one line
    weights = [min(40, max(7, round(1.25 * len(plain[0][c])), max(len(r[c]) for r in plain[1:])))
               for c in range(n_cols)]
    numeric_col = [all(NUMERIC.fullmatch(r[c]) for r in plain[1:]) for c in range(n_cols)]
    for c, wgt in enumerate(weights):
        tbl.columns[c].width = int(width * wgt / sum(weights))
    merge_first = n_rows > 12 and len({r[0] for r in plain[1:]}) < n_rows - 1  # A1: block column
    row_h = Pt(size * (2.5 if n_rows <= 6 else 2.0 if n_rows <= 12 else 1.45))  # 24-row ablation fits 4.7 in
    for r, row in enumerate(rows):
        tbl.rows[r].height = row_h
        for c, text in enumerate(row):
            cell = tbl.cell(r, c)
            if merge_first and c == 0 and r > 1 and plain[r][0] == plain[r - 1][0]:
                text = ""  # spanned by the merged block cell above
            right = c > 0 and numeric_col[c]  # numbers right-aligned, header included
            _paragraphs(cell.text_frame, [text], size, bullets=False,
                        color=ps.SURFACE if r == 0 else ps.INK,
                        align=PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT)
            if r == 0:
                for run in cell.text_frame.paragraphs[0].runs:
                    run.font.bold = True
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.margin_top = cell.margin_bottom = Inches(0.02 if n_rows <= 12 else 0.005)
    if merge_first:
        r = 1
        while r < n_rows:
            end = r
            while end + 1 < n_rows and plain[end + 1][0] == plain[r][0]:
                end += 1
            if end > r:
                tbl.cell(r, 0).merge(tbl.cell(end, 0))
            r = end + 1
    frame.left, frame.top, frame.width = left, top, width
    frame.height = row_h * n_rows


# ----------------------------------------------------------------------------- native diagrams
def _card(slide, origin, x: float, y: float, w: float, h: float, lines: list[str], edge: str,
          size: float = 14, sub_size: float = 12, bold_first: bool = False) -> None:
    """Rounded box in the figures' style: plum fill, colored edge, white text (first line bold)."""
    ox, oy = origin
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(ox + x), Inches(oy + y),
                                   Inches(w), Inches(h))
    shape.adjustments[0] = 0.1
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(ps.CARD)
    shape.line.color.rgb = _rgb(edge)
    shape.line.width = Pt(2)
    shape.shadow.inherit = False
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.clear()
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        main = i == 0 or not bold_first
        _runs(p, line, size if main else sub_size, color=ps.INK if main else ps.INK_2,
              bold=bold_first and i == 0, font="+mn-lt")


def _arrow(slide, origin, points: list[tuple[float, float]]) -> None:
    """Arrow along points (inches from origin); straight connector or an open polyline."""
    ox, oy = origin
    pts = [(ox + x, oy + y) for x, y in points]
    if len(pts) == 2:
        from pptx.enum.shapes import MSO_CONNECTOR

        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(pts[0][0]), Inches(pts[0][1]),
                                           Inches(pts[1][0]), Inches(pts[1][1]))
    else:
        # local unit = 0.01 in: python-pptx rounds local coordinates to integers, so inches would snap
        builder = slide.shapes.build_freeform(pts[0][0] * 100, pts[0][1] * 100, scale=Inches(0.01))
        builder.add_line_segments([(x * 100, y * 100) for x, y in pts[1:]], close=False)
        shape = builder.convert_to_shape()
        shape.fill.background()
        shape.shadow.inherit = False
    shape.line.color.rgb = _rgb(ps.INK_2)
    shape.line.width = Pt(1.75)
    ln = shape.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")


def _label(slide, origin, x: float, y: float, w: float, h: float, text: str, size: float = 13,
           color: str = ps.INK_2, bold: bool = False, align=PP_ALIGN.LEFT, rotation: float = 0) -> None:
    """Diagram annotation (no placeholder exists for it)."""
    ox, oy = origin
    box = slide.shapes.add_textbox(Inches(ox + x), Inches(oy + y), Inches(w), Inches(h))
    box.rotation = rotation
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _paragraphs(tf, text.split("\n"), size, bullets=False, color=color, align=align, font="+mn-lt")
    if bold:
        for p in tf.paragraphs:
            for run in p.runs:
                run.font.bold = True


def diagram_pipeline(slide, box, cfg: dict[str, Any]) -> None:
    """Slide 5: two preprocessing branches, two models, the blend, thresholds and rules."""
    o = (box[0] / EMU_PER_INCH, box[1] / EMU_PER_INCH)
    w = float(cfg["selected_model"]["blend"]["weight_on_finetuned_head"])
    _label(slide, o, 1.85, 0.0, 8.0, 0.32, "Rama transformer: texto casi original, con acentos y mayúsculas")
    _card(slide, o, 1.85, 0.38, 2.0, 0.95, ["Arreglo de codificación", "y espacios"], ps.TEAL)
    _card(slide, o, 4.2, 0.38, 4.15, 0.95, ["e5 multilingüe con fine-tuning", "(etapa 1): 10 probabilidades"], ps.TEAL)
    _card(slide, o, 0.0, 1.78, 1.45, 1.1, ["Respuesta", "en texto libre"], ps.INK_2)
    _card(slide, o, 1.85, 3.33, 2.0, 0.95, ["Lemas, sin acentos,", "stopwords propias"], ps.ROSE)
    _card(slide, o, 4.2, 3.33, 1.95, 0.95, ["TF-IDF palabras 1-2", "y caracteres 3-5"], ps.ROSE)
    _card(slide, o, 6.4, 3.33, 1.95, 0.95, ["Regresión logística", "por etiqueta (B2)"], ps.ROSE)
    _label(slide, o, 1.85, 4.36, 8.0, 0.32, "Rama léxica: texto lematizado y normalizado")
    _card(slide, o, 8.7, 1.83, 1.25, 1.0, ["Mezcla", f"{w:.1f} / {1 - w:.1f}"], ps.INK_2)
    _card(slide, o, 10.25, 1.13, 1.88, 2.4, ["Umbral por etiqueta", "(CV out-of-fold)", "",
                                            "Reglas: al menos", "una etiqueta;", "no excluyente"], ps.INK_2)
    for pts in ([(0.725, 1.78), (0.725, 0.855), (1.85, 0.855)], [(0.725, 2.88), (0.725, 3.805), (1.85, 3.805)],
                [(3.85, 0.855), (4.2, 0.855)], [(3.85, 3.805), (4.2, 3.805)], [(6.15, 3.805), (6.4, 3.805)],
                [(8.35, 0.855), (9.325, 0.855), (9.325, 1.83)], [(8.35, 3.805), (9.325, 3.805), (9.325, 2.83)],
                [(9.95, 2.33), (10.25, 2.33)], [(11.19, 3.53), (11.19, 3.93)]):
        _arrow(slide, o, pts)
    _label(slide, o, 10.25, 3.97, 1.88, 0.35, "etiquetas", size=15, color=ps.INK, bold=True, align=PP_ALIGN.CENTER)


def diagram_deployment(slide, box, cfg: dict[str, Any]) -> None:
    """Slide 10, left column: API, model, review routing, monitoring and the retraining loop."""
    import pandas as pd

    o = (box[0] / EMU_PER_INCH, box[1] / EMU_PER_INCH)
    lat = pd.read_csv(resolve(cfg["paths"]["latency_csv"])).set_index("model")
    e9, b2 = lat.loc["E9b (selected)"], lat.loc["B2 (fallback)"]
    margin, review = cfg["serving"]["review_margin"], " / ".join(cfg["serving"]["review_labels"])
    # rows end 3.7 in below the column top (6.0 in on the slide), above the layout's cone at 6.05 in
    _card(slide, o, 0.0, 0.0, 1.6, 0.85, ["POST /predict", "texto"], ps.INK_2, bold_first=True)
    _card(slide, o, 1.95, 0.0, 3.6, 0.85, ["E9b congelado en CPU", f"p50 {e9.p50_ms} ms, p95 {e9.p95_ms} ms",
                                          f"B2 de respaldo: {b2.p50_ms} ms"], ps.TEAL, bold_first=True)
    _card(slide, o, 0.0, 1.1, 5.55, 0.5, ["**Respuesta:** etiquetas, probabilidades y needs_review"], ps.INK_2,
          size=13)
    _card(slide, o, 0.0, 1.85, 2.65, 0.8, ["Etiqueta automática", "confiable y sin riesgo"], ps.TEAL, bold_first=True)
    _card(slide, o, 2.9, 1.85, 2.65, 0.8, ["Revisión humana", f"a {margin} del umbral, o {review}"], ps.ROSE,
          bold_first=True)
    _card(slide, o, 0.0, 2.95, 2.65, 0.75, ["Monitoreo", "mezcla de etiquetas y deriva"], ps.INK_2, bold_first=True)
    _card(slide, o, 2.9, 2.95, 2.65, 0.75, ["Reentrenamiento", "mismo protocolo de CV"], ps.TEAL, bold_first=True)
    for pts in ([(1.6, 0.425), (1.95, 0.425)], [(3.75, 0.85), (3.75, 1.1)], [(1.325, 1.6), (1.325, 1.85)],
                [(4.225, 1.6), (4.225, 1.85)], [(1.325, 2.65), (1.325, 2.95)], [(4.225, 2.65), (4.225, 2.95)],
                [(2.65, 3.325), (2.9, 3.325)], [(5.55, 3.325), (5.78, 3.325), (5.78, 0.425), (5.55, 0.425)]):
        _arrow(slide, o, pts)
    _label(slide, o, 4.3, 2.66, 1.25, 0.28, "etiquetas revisadas", size=10)
    # rotated label beside the loop's vertical segment; a rotated box keeps its unrotated left/top
    _label(slide, o, 5.93 - 0.8, 1.875 - 0.13, 1.6, 0.26, "nuevo modelo", size=11, align=PP_ALIGN.CENTER,
           rotation=270)


DIAGRAMS = {"pipeline": diagram_pipeline, "despliegue": diagram_deployment}


# ----------------------------------------------------------------------------- slides
def _drop_sample_slides(prs) -> None:
    """Remove the template's own slides; masters, layouts and theme stay."""
    ids = prs.slides._sldIdLst
    for sld_id in list(ids):
        prs.part.drop_rel(sld_id.rId)
        ids.remove(sld_id)


def _table_style_id(prs) -> str:
    """The template's default table style (its sample tables use it): teal header and rules."""
    part = prs.part.part_related_by(RT.TABLE_STYLES)
    return etree.fromstring(part.blob).get("def")


def _slide_number(slide) -> None:
    """Clone the layout's slide-number placeholder; python-pptx skips it when adding a slide."""
    for ph in slide.slide_layout.placeholders:
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER:
            el = copy.deepcopy(ph._element)
            ids = [int(i) for i in slide.shapes._spTree.xpath(".//p:cNvPr/@id")]
            el.find(".//" + qn("p:cNvPr")).set("id", str(max(ids, default=1) + 1))  # ids must be unique
            slide.shapes._spTree.append(el)


def _figure_path(cfg: dict[str, Any], name: str) -> Path:
    path = resolve(cfg["deck"]["figures_dir"]) / name
    if not path.exists():
        raise FileNotFoundError(f"{path}: run python -m intent.deck_figures (or app/screenshot.py)")
    return path


def _place(slide, kind: str, value, ph, box, cfg, style_id: str, caption: list[str] | None = None,
           tag: str | None = None) -> None:
    """Put one item (figure, table or native diagram) in a placeholder's area.

    With a caption, the picture goes at the top of the area and the placeholder moves below
    it to hold the caption, so text still lives in a placeholder rather than a free text box.
    """
    dpi = cfg["deck"]["dpi"]
    reserve = Inches(0.62 if tag else 0) + (Inches(0.25 + 0.34 * len(caption)) if caption else 0)
    inner = (box[0], box[1], box[2], box[3] - reserve)
    if kind == "diagram":
        _remove(ph)
        DIAGRAMS[value](slide, box, cfg)
    elif kind == "table":
        _table(slide, value, box, ph, style_id)
    elif caption:
        path = _figure_path(cfg, value)
        pic = _picture(slide, path, inner, None, dpi)
        top = pic.top + pic.height + Inches(0.08)
        _set_box(ph, box[0], top, box[2], box[1] + box[3] - top)
        ph.text_frame.vertical_anchor = MSO_ANCHOR.TOP
        _paragraphs(ph.text_frame, caption, CAPTION_PT, bullets=len(caption) > 1)
    else:
        path = _figure_path(cfg, value)
        screenshot = value.startswith("demo")
        pic = _picture(slide, path, inner, ph, None if screenshot else dpi)
        if screenshot:  # a light screenshot on the navy slide reads as a window with a hairline frame
            pic.line.color.rgb = _rgb(ps.INK_2)
            pic.line.width = Pt(0.75)
        if tag:
            tw = Inches(5.3)
            left = box[0] + (box[2] - tw) // 2
            top = pic.top + pic.height + Inches(0.14)
            _card(slide, (left / EMU_PER_INCH, top / EMU_PER_INCH), 0, 0, tw / EMU_PER_INCH, 0.44,
                  [f"**{tag.split(':')[0]}:**{tag.split(':', 1)[1]}" if ":" in tag else tag], ps.TEAL, size=14)


def _set_box(shape, left: int, top: int, width: int, height: int) -> None:
    """Set all four: a placeholder given only some of them loses the inherited rest."""
    shape.left, shape.top, shape.width, shape.height = int(left), int(top), int(width), int(height)


def _clear_of_decorations(slide, shape, margin: int = Inches(0.1)) -> None:
    """Shrink a text shape that overlaps one of the layout's decorative shapes.

    Layout decorations (spheres, cones) are drawn under the slide's shapes, so text that
    runs over them stays legible but looks careless. Prefer trimming the bottom (text
    flows from the top) while 60% of the height is left; otherwise trim the side the
    decoration is on.
    """
    for deco in slide.slide_layout.shapes:
        if deco.is_placeholder:
            continue
        l, t, w, h = _box(shape)
        r, b = l + w, t + h
        dl, dt, dr, db = deco.left, deco.top, deco.left + deco.width, deco.top + deco.height
        if dr <= l or dl >= r or db <= t or dt >= b:
            continue
        if dt - margin - t >= 0.6 * h:
            _set_box(shape, l, t, w, dt - margin - t)
        elif (dl + dr) / 2 > (l + r) / 2:
            _set_box(shape, l, t, dl - margin - l, h)
        else:
            _set_box(shape, dr + margin, t, r - dr - margin, h)


def _native_height(cfg: dict[str, Any], name: str) -> int:
    """Height in EMU at which a deck figure is inserted (its native size at the deck dpi)."""
    return Inches(Image.open(_figure_path(cfg, name)).size[1] / cfg["deck"]["dpi"])


def _captions(spec: SlideSpec, items, cols, cfg) -> list[list[str]]:
    """Text lines of a two-column slide, as captions under the figures that leave room for one.

    One line per such column when the counts match (slide 7); otherwise all lines go under
    the last column with room (slide 8, where the left figure fills its column).
    """
    room = [kind == "figure" and not value.startswith("demo")
            and _native_height(cfg, value) + Inches(0.55) <= ph.height
            for ph, (kind, value) in zip(cols, items)]
    out: list[list[str]] = [[] for _ in items]
    if not spec.lines:
        return out
    with_room = [i for i, ok in enumerate(room) if ok]
    if not with_room:
        raise ValueError(f"slide {spec.key}: no column has room for the text lines")
    if len(with_room) == len(spec.lines):
        for i, line in zip(with_room, spec.lines):
            out[i] = [line]
    else:
        out[with_room[-1]] = list(spec.lines)
    return out


def _items(spec: SlideSpec) -> list[tuple[str, Any]]:
    """Visual items in reading order (left column first)."""
    items: list[tuple[str, Any]] = []
    if spec.diagram:
        items.append(("diagram", spec.diagram))
    items += [("table", t) for t in spec.tables]
    items += [("figure", f) for f in spec.figures]
    return items


def build_slide(prs, spec: SlideSpec, cfg: dict[str, Any], style_id: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[spec.layout])
    phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
    _paragraphs(slide.shapes.title.text_frame, [spec.title], TITLE_APPENDIX_PT if spec.appendix else TITLE_PT,
                bullets=False)
    items = _items(spec)

    if spec.layout == 8:    # text left, table or figure right
        _paragraphs(phs[1].text_frame, spec.lines, BODY_PT)
        kind, value = items[0]
        if kind == "table":
            _table(slide, value, _box(phs[13]), phs[13], style_id)
        else:  # the right area is a table placeholder; a picture cannot live in it
            box = _box(phs[13])
            _remove(phs[13])
            _picture(slide, _figure_path(cfg, value), box, None, cfg["deck"]["dpi"])
    elif spec.layout == 10:  # one full-width item
        kind, value = items[0]
        _place(slide, kind, value, phs[1], _box(phs[1]), cfg, style_id)
    elif spec.layout == 6:   # two columns; text lines become captions
        cols = [phs[1], phs[13]]
        captions = _captions(spec, items, cols, cfg)
        for ph, (kind, value), caption in zip(cols, items, captions):
            _place(slide, kind, value, ph, _box(ph), cfg, style_id, caption or None, spec.tag if kind == "figure" else None)
    elif spec.layout == 4:   # text only
        _paragraphs(phs[1].text_frame, spec.lines, 24)
    else:
        raise ValueError(f"slide {spec.key}: layout {spec.layout} is not used by the outline")

    for ph in list(slide.placeholders):  # nothing may show "click to add text" in edit mode
        if ph.has_text_frame and not ph.text_frame.text.strip() and ph.placeholder_format.type != PP_PLACEHOLDER.TITLE:
            _remove(ph)
        elif ph.has_text_frame:
            _clear_of_decorations(slide, ph)
    _slide_number(slide)
    slide.notes_slide.notes_text_frame.text = _plain(spec.notes)


def build_deck(cfg: dict[str, Any] | None = None) -> Path:
    cfg = cfg or load_config()
    d = cfg["deck"]
    prs = Presentation(resolve(d["template"]))
    _drop_sample_slides(prs)
    style_id = _table_style_id(prs)
    for spec in parse_outline(resolve(d["outline"])):
        build_slide(prs, spec, cfg, style_id)
    out = resolve(d["output"])
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


# ----------------------------------------------------------------------------- language check
# Technical terms kept in English on purpose (reports/glossary_es.md, "sin traducir"), plus
# label codes and identifiers. Anything else that looks English is reported.
KEPT_IN_ENGLISH = {
    "macro", "micro", "samples", "hamming", "loss", "pr-auc", "mean", "tf-idf", "fine-tuning", "embedding",
    "embeddings", "features", "fold", "folds", "out-of-fold", "test", "baseline", "codebook", "stopwords",
    "mojibake", "ftfy", "transformer", "encoder", "lightgbm", "cleanlab", "needs_review", "post", "predict",
    "api", "cpu", "gpu", "recall", "kappa", "bootstrap", "commit", "tag", "hash", "config", "demo", "marketing",
    "wilson", "make", "csv", "llm", "lr", "mpnet", "pyme", "b2_margin", "b2_codebook",
    "crec", "cred", "equ", "inic", "inv", "mkt", "no", "renta", "sueldo", "temp",  # label codes
}
ENGLISH_WORDS = {
    "the", "and", "of", "to", "in", "is", "are", "for", "with", "on", "by", "from", "this", "that", "these",
    "be", "been", "as", "at", "or", "an", "it", "its", "not", "but", "which", "into", "than", "then", "when",
    "where", "after", "before", "only", "also", "all", "each", "both", "row", "rows", "label", "labels",
    "model", "models", "selected", "reference", "split", "grouped", "run", "results", "errors", "threshold",
    "thresholds", "rules", "rule", "blend", "frozen", "chain", "stage", "review", "human", "latency",
    "monitoring", "retraining", "payroll", "rent", "growth", "debts", "equipment", "inventory", "seasonal",
    "sales", "business", "use", "data", "set", "per", "multiple", "intents", "other", "short", "vague",
    "likely", "none", "slide", "notes", "figure", "table", "see", "how", "why", "what", "we", "our", "you",
    "will", "would", "should", "can", "could", "has", "have", "was", "were", "more", "most", "less", "best",
    "same", "new", "fixed", "cutoff", "close", "risk", "head", "weight", "group", "groups", "fitting",
}
ENGLISH_SUFFIX = re.compile(r"[a-z](ing|tion|ness|ly|ed)$")
DASHES = {"—": "em dash", "–": "en dash"}


def deck_text(path: Path) -> list[tuple[str, str]]:
    """(where, text) for every visible text run's paragraph, table cell and speaker note."""
    prs = Presentation(path)
    out = []

    def walk(shapes, where):
        for shape in shapes:
            if shape.shape_type == 6:  # group
                walk(shape.shapes, where)
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append((where, shape.text_frame.text))
            if getattr(shape, "has_table", False) and shape.has_table:
                out.extend((where + " tabla", cell.text) for row in shape.table.rows for cell in row.cells
                           if cell.text.strip())

    for i, slide in enumerate(prs.slides, start=1):
        walk(slide.shapes, f"diapositiva {i}")
        if slide.has_notes_slide:
            out.append((f"diapositiva {i} notas", slide.notes_slide.notes_text_frame.text))
    return out


def check_language(cfg: dict[str, Any] | None = None) -> str:
    """Report dashes and English words in the deck's visible text and notes ("ok" if none)."""
    cfg = cfg or load_config()
    problems = []
    for where, text in deck_text(resolve(cfg["deck"]["output"])):
        problems += [f"{where}: {name} in {text[:60]!r}" for ch, name in DASHES.items() if ch in text]
        for token in re.findall(r"[A-Za-zÁÉÍÓÚÜÑáéíóúüñ_][\w\-]*", text):
            t = token.lower()
            if t in KEPT_IN_ENGLISH or re.fullmatch(r"[bea]\d+\w*", t):  # experiment names B2, E9b, A1
                continue
            if t in ENGLISH_WORDS or (ENGLISH_SUFFIX.search(t) and not re.search(r"[áéíóúñ]", t)):
                problems.append(f"{where}: possible English {token!r}")
    return "ok" if not problems else "\n".join(sorted(set(problems)))


# ----------------------------------------------------------------------------- PDF and PNG
def _soffice(cfg: dict[str, Any]) -> str:
    for candidate in cfg["deck"]["soffice"]:
        found = shutil.which(candidate) or (candidate if Path(candidate).exists() else None)
        if found:
            return found
    raise FileNotFoundError(f"LibreOffice not found; tried {cfg['deck']['soffice']}")


def _lo_profile(root: Path, substitutes: dict[str, str]) -> None:
    """Throwaway LibreOffice profile whose font replacement table maps the cloud heading font.

    A separate profile keeps the user's own LibreOffice settings untouched.
    """
    pairs = "".join(
        f'<node oor:name="_{i}" oor:op="replace">'
        f'<prop oor:name="Always" oor:op="fuse"><value>true</value></prop>'
        f'<prop oor:name="OnScreenOnly" oor:op="fuse"><value>false</value></prop>'
        f'<prop oor:name="ReplaceFont" oor:op="fuse"><value>{src}</value></prop>'
        f'<prop oor:name="SubstituteFont" oor:op="fuse"><value>{dst}</value></prop></node>'
        for i, (src, dst) in enumerate(substitutes.items()))
    xcu = ('<?xml version="1.0" encoding="UTF-8"?>'
           '<oor:items xmlns:oor="http://openoffice.org/2001/registry" '
           'xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">'
           '<item oor:path="/org.openoffice.Office.Common/Font/Substitution">'
           '<prop oor:name="Replacement" oor:op="fuse"><value>true</value></prop></item>'
           f'<item oor:path="/org.openoffice.Office.Common/Font/Substitution/FontPairs">{pairs}</item>'
           '</oor:items>')
    (root / "user").mkdir(parents=True, exist_ok=True)
    (root / "user" / "registrymodifications.xcu").write_text(xcu, encoding="utf-8")


def export_pdf(cfg: dict[str, Any] | None = None) -> Path:
    """PDF copy of the deck with LibreOffice headless."""
    cfg = cfg or load_config()
    pptx = resolve(cfg["deck"]["output"])
    with tempfile.TemporaryDirectory() as profile:
        _lo_profile(Path(profile), cfg["deck"]["pdf_font_substitutes"])
        subprocess.run([_soffice(cfg), f"-env:UserInstallation={Path(profile).as_uri()}", "--headless",
                        "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                       check=True, timeout=600, capture_output=True)
    return pptx.with_suffix(".pdf")


def render_png(cfg: dict[str, Any] | None = None) -> list[Path]:
    """One PNG per page of the PDF copy, for visual review (old renders are replaced)."""
    import pymupdf

    cfg = cfg or load_config()
    out = resolve(cfg["deck"]["png_dir"])
    out.mkdir(parents=True, exist_ok=True)
    for old in out.glob("slide_*.png"):
        old.unlink()
    paths = []
    with pymupdf.open(resolve(cfg["deck"]["output"]).with_suffix(".pdf")) as doc:
        for i, page in enumerate(doc, start=1):
            path = out / f"slide_{i:02d}.png"
            page.get_pixmap(dpi=cfg["deck"]["png_dpi"]).save(path)
            paths.append(path)
    return paths


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    steps = {"inspect": write_inspection, "build": build_deck, "pdf": export_pdf, "png": render_png,
             "check": check_language}
    args = sys.argv[1:]
    if args == ["all"]:
        args = ["build", "pdf", "png", "check"]
    if not args or any(a not in steps for a in args):
        sys.exit("usage: python -m intent.deck inspect | build | pdf | png | check | all")
    for a in args:
        print(steps[a]())
